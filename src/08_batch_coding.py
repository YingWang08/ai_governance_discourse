#!/usr/bin/env python3
"""
08_batch_coding.py  —  批量编码 / 批量裁决（替代 06 的逐句交互）
================================================================
解决"一个一个回车太慢"的问题：把所有待编码 / 待裁决的句子一次性
导出成 Excel，两位编码者直接在下拉框里填，填完一次性导回。

与 06_interactive_kappa.py 完全兼容：读写同一个 coding_sample.csv，
标签词表 {cap, con, both, neither}，金标准口径一致。

典型流程
--------
  # 1. 先用 06 抽样（可把 n 放大到 300/500 来收窄置信区间）
  python 06_interactive_kappa.py --make-sample --n 300 --seed 99

  # 2. 导出空白编码表（默认一份工作簿，内含 Coder_1 / Coder_2 两个表）
  python 08_batch_coding.py --export-coding
  #    两位编码者各自在 Coder_1 / Coder_2 表的 label 列下拉填写（互不参看）

  # 3. 填完导回 coding_sample.csv
  python 08_batch_coding.py --import-coding --in coding_blank.xlsx

  # 4. 出 Kappa（用 06 现成的打分，含 CI / alpha / 逐类别）
  python 06_interactive_kappa.py --score

  # 5. 一次性导出所有分歧，批量裁决
  python 08_batch_coding.py --export-adjudication
  #    在 adjudicated 列下拉给每条分歧定一个最终标签

  # 6. 导回裁决结果，再 score 一次得到金标准
  python 08_batch_coding.py --import-adjudication --in adjudication.xlsx
  python 06_interactive_kappa.py --score
"""
from __future__ import annotations
import os, re, sys, argparse
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE) if os.path.basename(HERE) == "src" else HERE
PROC   = os.path.join(ROOT, "data", "processed")
CONF   = os.path.join(ROOT, "config")
TABLES = os.path.join(ROOT, "output", "tables")
SAMPLE = os.path.join(TABLES, "coding_sample.csv")

VALID = ["cap", "con", "both", "neither"]
LABEL_CN = {"cap": "能力", "con": "后果", "both": "兼有", "neither": "均无"}


# ── 词典命中（仅用于裁决表里给编码者参考，逻辑与 06 一致）──────────────
def _load_lexicon():
    import yaml
    p = os.path.join(CONF, "lexicon_expanded.yaml")
    if not os.path.exists(p):
        p = os.path.join(CONF, "lexicon.yaml")
    if not os.path.exists(p):
        return {"capability": [], "consequence": [], "boundary": []}
    with open(p, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def _matcher(terms):
    if not terms:
        return None
    esc = sorted([re.escape(t.lower()) for t in terms], key=len, reverse=True)
    return re.compile(r"\b(?:" + "|".join(esc) + r")\b", re.I)


def _hits(text, m):
    if m is None:
        return []
    return sorted(set(x.lower() for x in m.findall(str(text))))


# ── openpyxl 下拉框辅助 ────────────────────────────────────────────────
def _add_dropdown(ws, col_letter, n_rows, choices):
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"',
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{n_rows + 1}")


def _autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _read_sample():
    if not os.path.exists(SAMPLE):
        sys.exit(f"[!] 找不到 {SAMPLE}，请先运行 06_interactive_kappa.py --make-sample")
    return pd.read_csv(SAMPLE, dtype=str).fillna("")


# ═══════════════════════════════════════════════════════════════════════
# 导出 / 导入：盲编
# ═══════════════════════════════════════════════════════════════════════
def export_coding(out_path, only_coder=None):
    from openpyxl import Workbook
    df = _read_sample()
    has_cn = "text_cn" in df.columns
    wb = Workbook()

    # 说明页
    info = wb.active
    info.title = "Instructions"
    info.append(["批量编码说明 / Coding instructions"])
    info.append([""])
    info.append(["在 label 列的下拉框中为每一句选择一个标签："])
    info.append(["cap = 能力框架 (capability)"])
    info.append(["con = 后果框架 (consequence)"])
    info.append(["both = 两者皆有"])
    info.append(["neither = 两者皆无（含程序性/定义性条款）"])
    info.append([""])
    info.append(["★ 两位编码者请各自填写自己的工作表，填写期间不要查看对方的表。"])
    info.append(["★ 仅依据句子本身判断；本表不显示自动标注，以保证盲编。"])
    _autosize(info, {"A": 70})

    coders = [only_coder] if only_coder else [1, 2]
    for c in coders:
        ws = wb.create_sheet(title=f"Coder_{c}")
        headers = ["sent_id", "year", "text"] + (["text_cn"] if has_cn else []) + ["label"]
        ws.append(headers)
        for _, r in df.iterrows():
            row = [r["sent_id"], r.get("year", ""), r["text"]]
            if has_cn:
                row.append(r.get("text_cn", ""))
            row.append("")  # 空 label
            ws.append(row)
        label_col = chr(ord("A") + len(headers) - 1)  # 最后一列
        _add_dropdown(ws, label_col, len(df), VALID)
        widths = {"A": 12, "B": 8, "C": 90}
        if has_cn:
            widths["D"] = 60
        widths[label_col] = 12
        _autosize(ws, widths)
        ws.freeze_panes = "A2"

    wb.save(out_path)
    print(f"[\u2713] 已导出空白编码表 \u2192 {out_path}")
    print(f"    句子数 {len(df)}；在 Coder_1 / Coder_2 表的 label 列下拉填写后，")
    print(f"    运行：python 08_batch_coding.py --import-coding --in {os.path.basename(out_path)}")


def import_coding(in_path):
    from openpyxl import load_workbook
    if not os.path.exists(in_path):
        sys.exit(f"[!] 找不到 {in_path}")
    df = _read_sample().set_index("sent_id")
    wb = load_workbook(in_path, data_only=True)
    target = {1: "coder_1_blind", 2: "coder_2_blind", 3: "coder_3_blind"}
    n_written = {}
    for c, col in target.items():
        sheet = f"Coder_{c}"
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        try:
            i_id = header.index("sent_id"); i_lab = header.index("label")
        except ValueError:
            sys.exit(f"[!] 工作表 {sheet} 需包含 sent_id 与 label 列")
        if col not in df.columns:
            df[col] = ""
        cnt = 0
        for row in rows[1:]:
            sid = row[i_id]
            lab = (str(row[i_lab]).strip().lower() if row[i_lab] is not None else "")
            if sid is None:
                continue
            sid = str(sid)
            if lab and lab not in VALID:
                print(f"    [warn] sent_id={sid} 标签 '{lab}' 不在 {VALID}，已跳过")
                continue
            if sid in df.index and lab:
                df.at[sid, col] = lab; cnt += 1
        n_written[sheet] = cnt
    df.reset_index().to_csv(SAMPLE, index=False)
    print(f"[\u2713] 已写回 {SAMPLE}")
    for k, v in n_written.items():
        print(f"    {k}: {v} 条")
    print("    下一步：python 06_interactive_kappa.py --score")


# ═══════════════════════════════════════════════════════════════════════
# 导出 / 导入：批量裁决
# ═══════════════════════════════════════════════════════════════════════
def export_adjudication(out_path):
    from openpyxl import Workbook
    df = _read_sample()
    blind = df[df.coder_1_blind.isin(VALID) & df.coder_2_blind.isin(VALID)].copy()
    if blind.empty:
        sys.exit("[!] 还没有完成盲编，无法导出分歧。请先 --import-coding。")
    dis = blind[blind.coder_1_blind != blind.coder_2_blind].copy()
    if dis.empty:
        print("[\u2713] 没有分歧需要裁决，金标准 = 盲编一致标签。直接跑 06 --score 即可。")
        return

    lex = _load_lexicon()
    cap_m, con_m = _matcher(lex.get("capability", [])), _matcher(lex.get("consequence", []))
    has_cn = "text_cn" in dis.columns

    def suggest(r):
        c1, c2, a = r.coder_1_blind, r.coder_2_blind, r.auto_label
        if c1 == a: return a
        if c2 == a: return a
        return a  # 两人都与自动不一致时默认自动标注，供编码者参考
    dis["suggested"] = dis.apply(suggest, axis=1)

    wb = Workbook(); ws = wb.active; ws.title = "Adjudication"
    headers = ["sent_id", "year", "text"] + (["text_cn"] if has_cn else []) + \
              ["coder_1", "coder_2", "auto", "cap_terms", "con_terms",
               "suggested", "adjudicated", "note"]
    ws.append(headers)
    for _, r in dis.iterrows():
        row = [r["sent_id"], r.get("year", ""), r["text"]]
        if has_cn:
            row.append(r.get("text_cn", ""))
        row += [r["coder_1_blind"], r["coder_2_blind"], r["auto_label"],
                ", ".join(_hits(r["text"], cap_m)), ", ".join(_hits(r["text"], con_m)),
                r["suggested"], "", ""]
        ws.append(row)
    adj_col = chr(ord("A") + headers.index("adjudicated"))
    _add_dropdown(ws, adj_col, len(dis), VALID)
    widths = {"A": 12, "B": 8, "C": 80}
    if has_cn: widths["D"] = 55
    _autosize(ws, widths)
    ws.freeze_panes = "A2"
    wb.save(out_path)
    print(f"[\u2713] 已导出 {len(dis)} 条分歧 \u2192 {out_path}")
    print(f"    在 adjudicated 列下拉为每条定一个最终标签（可参考 suggested / 命中词），")
    print(f"    然后：python 08_batch_coding.py --import-adjudication --in {os.path.basename(out_path)}")


def import_adjudication(in_path):
    from openpyxl import load_workbook
    if not os.path.exists(in_path):
        sys.exit(f"[!] 找不到 {in_path}")
    df = _read_sample().set_index("sent_id")
    for col in ("adjudicated", "adj_note"):
        if col not in df.columns:
            df[col] = ""
    wb = load_workbook(in_path, data_only=True)
    ws = wb["Adjudication"] if "Adjudication" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    i_id = header.index("sent_id"); i_adj = header.index("adjudicated")
    i_note = header.index("note") if "note" in header else None
    cnt = 0
    for row in rows[1:]:
        sid = row[i_id]
        if sid is None:
            continue
        sid = str(sid)
        lab = (str(row[i_adj]).strip().lower() if row[i_adj] is not None else "")
        if lab and lab not in VALID:
            print(f"    [warn] sent_id={sid} 裁决标签 '{lab}' 非法，已跳过"); continue
        if sid in df.index and lab:
            df.at[sid, "adjudicated"] = lab
            if i_note is not None and row[i_note]:
                df.at[sid, "adj_note"] = str(row[i_note])
            cnt += 1
    df.reset_index().to_csv(SAMPLE, index=False)
    print(f"[\u2713] 已写回 {cnt} 条裁决 \u2192 {SAMPLE}")
    print("    下一步：python 06_interactive_kappa.py --score  （此时含金标准）")


def main():
    ap = argparse.ArgumentParser(description="批量编码 / 批量裁决（Excel 一次性填写）")
    ap.add_argument("--export-coding", action="store_true")
    ap.add_argument("--import-coding", action="store_true")
    ap.add_argument("--export-adjudication", action="store_true")
    ap.add_argument("--import-adjudication", action="store_true")
    ap.add_argument("--coder", type=int, choices=[1, 2, 3], default=None,
                    help="仅导出单个编码者的表（默认两个都导）")
    ap.add_argument("--in", dest="in_path", default=None)
    ap.add_argument("--out", dest="out_path", default=None)
    a = ap.parse_args()

    if a.export_coding:
        out = a.out_path or "coding_blank.xlsx"
        export_coding(out, only_coder=a.coder)
    elif a.import_coding:
        import_coding(a.in_path or "coding_blank.xlsx")
    elif a.export_adjudication:
        export_adjudication(a.out_path or "adjudication.xlsx")
    elif a.import_adjudication:
        import_adjudication(a.in_path or "adjudication.xlsx")
    else:
        ap.print_help()


if __name__ == "__main__":
    main()
